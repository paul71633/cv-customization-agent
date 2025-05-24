import os
import re
import json
from typing import Dict, List, Optional, Set, Any
from collections import Counter
from io import BytesIO
import PyPDF2
from docx import Document

class CompleteCVAgent:
    """Complete CV Analysis and Job Matching Agent - Local Version"""

    def __init__(self):
        self.cv_results = None

        # Comprehensive skill database
        self.skill_categories = {
            'programming': ['python', 'java', 'javascript', 'typescript', 'c++', 'c#', 'ruby', 'php', 'go', 'rust', 'swift', 'kotlin', 'scala', 'r', 'matlab'],
            'web_frontend': ['html', 'css', 'react', 'angular', 'vue', 'jquery', 'bootstrap', 'sass', 'less', 'webpack', 'babel'],
            'web_backend': ['node.js', 'express', 'django', 'flask', 'spring', 'laravel', 'rails', 'asp.net', 'fastapi'],
            'database': ['sql', 'mysql', 'postgresql', 'mongodb', 'redis', 'oracle', 'sqlite', 'cassandra', 'dynamodb'],
            'cloud_devops': ['aws', 'azure', 'gcp', 'docker', 'kubernetes', 'terraform', 'jenkins', 'ansible', 'heroku', 'vercel'],
            'data_science': ['pandas', 'numpy', 'scikit-learn', 'tensorflow', 'pytorch', 'spark', 'hadoop', 'tableau', 'powerbi'],
            'mobile': ['android', 'ios', 'react native', 'flutter', 'xamarin', 'swift', 'kotlin'],
            'tools_other': ['git', 'jira', 'confluence', 'slack', 'trello', 'figma', 'photoshop', 'linux', 'windows', 'macos']
        }

        # Soft skills
        self.soft_skills = [
            'leadership', 'communication', 'teamwork', 'problem solving', 'analytical',
            'creative', 'adaptable', 'organized', 'detail oriented', 'time management',
            'project management', 'agile', 'scrum', 'mentoring', 'collaboration'
        ]

    # ==================== FILE HANDLING ====================

    def analyze_cv_from_file(self, file_path: str) -> Dict[str, Any]:
        """Analyze CV from local file path"""
        if not os.path.exists(file_path):
            raise FileNotFoundError(f"File not found: {file_path}")
        
        filename = os.path.basename(file_path)
        print(f"\n🔍 ANALYZING CV: {filename}")
        print("=" * 60)

        # Read file content
        with open(file_path, 'rb') as file:
            file_content = file.read()
        
        file_size = len(file_content) / 1024  # KB
        print(f"📊 File size: {file_size:.1f} KB")

        # Extract text based on file extension
        file_extension = filename.lower().split('.')[-1]
        
        if file_extension == 'pdf':
            text = self.extract_text_from_pdf_bytes(file_content)
        elif file_extension in ['docx', 'doc']:
            text = self.extract_text_from_docx_bytes(file_content)
        else:
            raise ValueError(f"Unsupported format: {file_extension}")

        return self._analyze_cv_text(text, filename)

    def extract_text_from_pdf_bytes(self, file_content: bytes) -> str:
        """Extract text from PDF bytes"""
        try:
            pdf_file = BytesIO(file_content)
            pdf_reader = PyPDF2.PdfReader(pdf_file)
            text = ""

            print(f"   📖 PDF has {len(pdf_reader.pages)} page(s)")

            for i, page in enumerate(pdf_reader.pages, 1):
                page_text = page.extract_text()
                text += page_text + "\n"
                print(f"   ✓ Page {i}: {len(page_text)} characters")

            return text.strip()
        except Exception as e:
            raise Exception(f"PDF reading error: {str(e)}")

    def extract_text_from_docx_bytes(self, file_content: bytes) -> str:
        """Extract text from DOCX bytes"""
        try:
            docx_file = BytesIO(file_content)
            doc = Document(docx_file)
            text = ""

            print(f"   📖 DOCX has {len(doc.paragraphs)} paragraph(s)")

            # Extract from paragraphs
            for paragraph in doc.paragraphs:
                if paragraph.text.strip():
                    text += paragraph.text + "\n"

            # Extract from tables
            table_count = 0
            for table in doc.tables:
                table_count += 1
                for row in table.rows:
                    for cell in row.cells:
                        if cell.text.strip():
                            text += cell.text + " "
                    text += "\n"

            if table_count > 0:
                print(f"   ✓ Also extracted from {table_count} table(s)")

            return text.strip()
        except Exception as e:
            raise Exception(f"DOCX reading error: {str(e)}")

    # ==================== CV ANALYSIS ====================

    def _analyze_cv_text(self, text: str, filename: str) -> Dict[str, Any]:
        """Internal method to analyze extracted CV text"""
        print(f"\n📊 EXTRACTION COMPLETE")
        print(f"   Total characters: {len(text):,}")
        print(f"   Total words: {len(text.split()):,}")

        # Analyze components
        print(f"\n📋 ANALYZING SECTIONS...")
        sections = self.parse_cv_sections(text)

        print(f"\n👤 EXTRACTING CONTACT INFO...")
        contact_details = self.extract_contact_details(text)

        print(f"\n🛠️ CATEGORIZING SKILLS...")
        skills_by_category = self.extract_skills_by_category(text)

        # Store results
        self.cv_results = {
            'filename': filename,
            'raw_text': text,
            'sections': sections,
            'skills_by_category': skills_by_category,
            'contact_details': contact_details,
            'statistics': {
                'word_count': len(text.split()),
                'char_count': len(text),
                'total_skills': sum(len(skills) for skills in skills_by_category.values()),
                'sections_found': len([s for s in sections.values() if s.strip()])
            }
        }

        # Print detailed summary
        self.print_cv_summary()
        return self.cv_results

    def parse_cv_sections(self, text: str) -> Dict[str, str]:
        """Parse CV into sections with enhanced detection"""
        sections = {
            'contact_info': '',
            'summary': '',
            'experience': '',
            'skills': '',
            'education': '',
            'certifications': '',
            'projects': '',
            'other': ''
        }

        section_patterns = {
            'contact_info': r'(contact|phone|email|address|linkedin|mobile|cell)',
            'summary': r'(summary|objective|profile|about|overview|professional summary)',
            'experience': r'(experience|work|employment|career|position|professional|history)',
            'skills': r'(skills|competencies|expertise|technologies|technical|proficiencies)',
            'education': r'(education|academic|degree|university|college|school|qualification)',
            'certifications': r'(certification|certificate|license|credential)',
            'projects': r'(projects|portfolio|work samples|personal projects)'
        }

        lines = text.split('\n')
        current_section = 'other'

        for line in lines:
            line = line.strip()
            if not line:
                continue

            # Check for section headers (short lines that match patterns)
            if len(line) < 100:
                line_lower = line.lower()
                for section, pattern in section_patterns.items():
                    if re.search(pattern, line_lower):
                        current_section = section
                        print(f"   🔍 Found section: {section.upper()}")
                        break

            # Add content to current section
            if sections[current_section]:
                sections[current_section] += "\n" + line
            else:
                sections[current_section] = line

        return sections

    def extract_contact_details(self, text: str) -> Dict[str, str]:
        """Extract contact information"""
        contact = {}

        # Email
        email_pattern = r'\b[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Z|a-z]{2,}\b'
        emails = re.findall(email_pattern, text)
        if emails:
            contact['email'] = emails[0]

        # Phone
        phone_pattern = r'(\+?1[-.\\s]?)?\\(?\\d{3}\\)?[-.\\s]?\\d{3}[-.\\s]?\\d{4}'
        phones = re.findall(phone_pattern, text)
        if phones:
            contact['phone'] = ''.join(phones[0]) if isinstance(phones[0], tuple) else phones[0]

        # LinkedIn
        linkedin_pattern = r'linkedin\.com/in/[\w-]+'
        linkedin = re.findall(linkedin_pattern, text.lower())
        if linkedin:
            contact['linkedin'] = linkedin[0]

        return contact

    def extract_skills_by_category(self, text: str) -> Dict[str, List[str]]:
        """Extract and categorize skills"""
        found_skills = {category: [] for category in self.skill_categories.keys()}
        found_skills['soft_skills'] = []

        text_lower = text.lower()

        # Technical skills
        for category, skills in self.skill_categories.items():
            for skill in skills:
                patterns = [
                    rf'\b{re.escape(skill.lower())}\b',
                    rf'{re.escape(skill.lower().replace(".", ""))}'
                ]

                for pattern in patterns:
                    if re.search(pattern, text_lower):
                        found_skills[category].append(skill)
                        break

        # Soft skills
        for skill in self.soft_skills:
            if re.search(rf'\b{re.escape(skill.lower())}\b', text_lower):
                found_skills['soft_skills'].append(skill)

        # Remove empty categories and duplicates
        return {k: list(set(v)) for k, v in found_skills.items() if v}

    # ==================== JOB ANALYSIS ====================

    def analyze_job_description(self, job_description: str) -> Dict[str, Any]:
        """Analyze job description for requirements"""
        skills_found = {'technical': [], 'soft': []}

        text_lower = job_description.lower()

        # Find technical skills
        all_tech_skills = []
        for category_skills in self.skill_categories.values():
            all_tech_skills.extend(category_skills)

        for skill in all_tech_skills:
            if re.search(rf'\b{re.escape(skill.lower())}\b', text_lower):
                skills_found['technical'].append(skill)

        # Find soft skills
        for skill in self.soft_skills:
            if re.search(rf'\b{re.escape(skill.lower())}\b', text_lower):
                skills_found['soft'].append(skill)

        # Extract experience requirements
        years_pattern = r'(\d+)(?:\+|\s*to\s*\d+)?\s*(?:years?|yrs?)\s*(?:of\s*)?(?:experience|exp)'
        years_matches = re.findall(years_pattern, text_lower)
        min_years = min([int(match) for match in years_matches]) if years_matches else None

        return {
            'required_skills': skills_found,
            'min_years_experience': min_years,
            'word_count': len(job_description.split())
        }

    def test_job_matching(self, job_description: str) -> Dict[str, Any]:
        """Test CV against job description with detailed output"""
        if not self.cv_results:
            print("❌ No CV analyzed yet. Analyze a CV first!")
            return None

        print(f"\n🎯 JOB MATCHING ANALYSIS")
        print("=" * 50)

        # Analyze job
        job_analysis = self.analyze_job_description(job_description)

        # Get your skills (flatten categories)
        your_skills = set()
        for category_skills in self.cv_results['skills_by_category'].values():
            your_skills.update([skill.lower() for skill in category_skills])

        # Get job requirements
        job_tech_skills = set([skill.lower() for skill in job_analysis['required_skills']['technical']])
        job_soft_skills = set([skill.lower() for skill in job_analysis['required_skills']['soft']])

        # Calculate matches
        tech_matches = your_skills.intersection(job_tech_skills)
        tech_missing = job_tech_skills - your_skills
        soft_matches = your_skills.intersection(job_soft_skills)
        soft_missing = job_soft_skills - your_skills

        # Print results
        print(f"🔧 TECHNICAL SKILLS:")
        print(f"   ✅ You HAVE: {sorted(list(tech_matches)) if tech_matches else 'None'}")
        print(f"   ❌ You're MISSING: {sorted(list(tech_missing)) if tech_missing else 'None'}")
        if job_tech_skills:
            tech_rate = len(tech_matches)/len(job_tech_skills)*100
            print(f"   📊 Technical Match Rate: {tech_rate:.1f}%")

        print(f"\n🤝 SOFT SKILLS:")
        print(f"   ✅ You HAVE: {sorted(list(soft_matches)) if soft_matches else 'None'}")
        print(f"   ❌ You're MISSING: {sorted(list(soft_missing)) if soft_missing else 'None'}")
        if job_soft_skills:
            soft_rate = len(soft_matches)/len(job_soft_skills)*100
            print(f"   📊 Soft Skills Match Rate: {soft_rate:.1f}%")

        # Overall match rate
        all_job_skills = job_tech_skills.union(job_soft_skills)
        all_matches = tech_matches.union(soft_matches)
        overall_rate = len(all_matches)/len(all_job_skills)*100 if all_job_skills else 0

        print(f"\n🎯 OVERALL MATCH RATE: {overall_rate:.1f}%")

        if job_analysis['min_years_experience']:
            print(f"📅 Experience Required: {job_analysis['min_years_experience']}+ years")

        return {
            'technical_matches': list(tech_matches),
            'technical_missing': list(tech_missing),
            'soft_matches': list(soft_matches),
            'soft_missing': list(soft_missing),
            'overall_match_rate': overall_rate,
            'min_years_required': job_analysis['min_years_experience']
        }

    def print_cv_summary(self):
        """Print detailed CV analysis summary"""
        if not self.cv_results:
            return

        print(f"\n" + "=" * 60)
        print("📋 COMPLETE CV ANALYSIS SUMMARY")
        print("=" * 60)

        # Contact details
        print(f"\n📞 CONTACT INFORMATION:")
        if self.cv_results['contact_details']:
            for key, value in self.cv_results['contact_details'].items():
                print(f"   ✓ {key.title()}: {value}")
        else:
            print("   ⚠️ No contact details auto-detected")

        # Sections
        print(f"\n📝 SECTIONS DETECTED:")
        for section, content in self.cv_results['sections'].items():
            if content.strip():
                word_count = len(content.split())
                preview = content[:80].replace('\n', ' ')
                print(f"   ✓ {section.title()}: {word_count} words")
                print(f"     Preview: '{preview}...'")

        # Skills by category
        print(f"\n🛠️ SKILLS BY CATEGORY:")
        total_skills = 0
        for category, skill_list in self.cv_results['skills_by_category'].items():
            if skill_list:
                print(f"   📌 {category.upper().replace('_', ' ')}: {len(skill_list)} skills")
                print(f"      → {', '.join(skill_list)}")
                total_skills += len(skill_list)

        # Statistics
        stats = self.cv_results['statistics']
        print(f"\n📈 OVERALL STATISTICS:")
        print(f"   📄 Total words: {stats['word_count']:,}")
        print(f"   🔤 Total characters: {stats['char_count']:,}")
        print(f"   🛠️ Total skills found: {stats['total_skills']}")
        print(f"   📋 Sections with content: {stats['sections_found']}")

        print(f"\n✅ CV analysis complete!")

    def save_results_to_json(self, output_path: str = None):
        """Save CV analysis results to JSON file"""
        if not self.cv_results:
            print("❌ No CV results to save!")
            return
        
        if not output_path:
            output_path = f"{self.cv_results['filename']}_analysis.json"
        
        try:
            with open(output_path, 'w', encoding='utf-8') as f:
                json.dump(self.cv_results, f, indent=2, ensure_ascii=False)
            print(f"✅ Results saved to: {output_path}")
        except Exception as e:
            print(f"❌ Error saving results: {e}")